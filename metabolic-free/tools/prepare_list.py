#!/usr/bin/env python3
"""Clean the warm contact list before the launch blast.

Reads a CSV (or .xlsx if openpyxl is installed), finds the name / email / phone
columns whatever they're called, then:
  - trims, lowercases and validates emails, removes duplicates
  - normalizes US phone numbers to +1XXXXXXXXXX
  - marks who can get texts: only rows with an SMS-consent column that says yes
  - writes files ready to import into the CRM

Outputs go to metabolic-free/data/private/ (git-ignored, never committed):
  email_list.csv        everyone with a valid email -> launch email sequence
  sms_consented.csv     valid phone + recorded SMS consent -> launch texts
  phone_no_consent.csv  phone but no recorded consent -> personal 1:1 texts or opt-in ask only
  rejected.csv          no usable email or phone, with the reason

Usage:  python3 tools/prepare_list.py path/to/contacts.csv
"""
import csv
import re
import sys
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "data" / "private"
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$")
ROLE_PREFIXES = ("info@", "admin@", "noreply@", "no-reply@", "support@", "sales@", "office@")
YES = {"yes", "y", "true", "1", "opted in", "opt-in", "subscribed"}

HEADER_HINTS = {
    "first_name": ("first name", "first_name", "firstname", "first"),
    "last_name": ("last name", "last_name", "lastname", "last", "surname"),
    "full_name": ("name", "full name", "full_name", "contact", "client"),
    "email": ("email", "e-mail", "email address", "mail"),
    "phone": ("phone", "mobile", "cell", "phone number", "telephone", "whatsapp"),
    "sms_consent": ("sms consent", "sms_consent", "text consent", "sms opt in", "sms opt-in", "texting"),
}


def read_rows(path):
    if path.suffix.lower() in (".xlsx", ".xls"):
        from openpyxl import load_workbook  # pip install openpyxl
        ws = load_workbook(path, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h or "").strip() for h in rows[0]]
        return header, [["" if v is None else str(v) for v in r] for r in rows[1:]]
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    return [h.strip() for h in rows[0]], rows[1:]


def map_columns(header):
    lower = [h.lower() for h in header]
    found = {}
    for field, hints in HEADER_HINTS.items():
        for hint in hints:
            if hint in lower and lower.index(hint) not in found.values():
                found[field] = lower.index(hint)
                break
    return found


def clean_phone(raw):
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) == 10:
        return "+1" + digits
    if len(digits) == 11 and digits.startswith("1"):
        return "+" + digits
    if len(digits) > 11 and (raw or "").strip().startswith("+"):
        return "+" + digits  # international, keep as given
    return ""


def main(path):
    header, rows = read_rows(Path(path))
    cols = map_columns(header)
    if "email" not in cols and "phone" not in cols:
        sys.exit(f"Couldn't find an email or phone column in: {header}")
    get = lambda r, k: r[cols[k]].strip() if k in cols and cols[k] < len(r) else ""

    seen, email_list, sms_ok, phone_only, rejected = set(), [], [], [], []
    for r in rows:
        first = get(r, "first_name") or get(r, "full_name").split(" ")[0]
        first = first.title()
        email = get(r, "email").lower()
        phone = clean_phone(get(r, "phone"))
        consent = get(r, "sms_consent").lower() in YES
        key = email or phone
        if not key:
            rejected.append({"row": r, "reason": "no email or phone"}); continue
        if key in seen:
            rejected.append({"row": r, "reason": "duplicate"}); continue
        seen.add(key)
        valid_email = bool(EMAIL_RE.match(email)) and not email.startswith(ROLE_PREFIXES)
        rec = {"first_name": first, "email": email if valid_email else "", "phone": phone,
               "tags": "warm_list_launch", "source": "gfc_xtreme_warm_list"}
        if valid_email:
            email_list.append(rec)
        if phone and consent:
            sms_ok.append(rec)
        elif phone:
            phone_only.append(rec)
        if not valid_email and not phone:
            rejected.append({"row": r, "reason": f"invalid email '{email}' and no phone"})

    OUT.mkdir(parents=True, exist_ok=True)
    fields = ["first_name", "email", "phone", "tags", "source"]
    for name, data in [("email_list", email_list), ("sms_consented", sms_ok), ("phone_no_consent", phone_only)]:
        with open(OUT / f"{name}.csv", "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(data)
    with open(OUT / "rejected.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["reason"] + header)
        for x in rejected:
            w.writerow([x["reason"]] + x["row"])

    print(f"Columns used: { {k: header[v] for k, v in cols.items()} }")
    print(f"Rows read ..................... {len(rows)}")
    print(f"Email list .................... {len(email_list)}")
    print(f"SMS with recorded consent ..... {len(sms_ok)}")
    print(f"Phone, no recorded consent .... {len(phone_only)}  -> 1:1 personal texts / opt-in ask only")
    print(f"Rejected ...................... {len(rejected)}")
    print(f"Files written to {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
